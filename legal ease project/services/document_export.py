import io
import os
import re
from typing import List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF


def sanitize_text(text: str) -> str:
    """
    Clean text for document export.
    """

    if not text:
        return ""

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2026": "...",
        "\u00a0": " ",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    return text.strip()


def split_document(text: str):
    """
    Separate the document into paragraphs/sections.
    """

    clean = sanitize_text(text)

    blocks = [
        block.strip()
        for block in clean.split("\n")
        if block.strip()
    ]

    return blocks


def format_docx(
    text: str,
    doc_type: str
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles

    normal_style = styles["Normal"]

    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(
        doc_type.upper()
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    document.add_paragraph()

    blocks = split_document(text)

    for block in blocks:

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(7)
        paragraph.paragraph_format.line_spacing = 1.15

        if is_heading(block):

            run = paragraph.add_run(block)

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        else:

            run = paragraph.add_run(block)

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    footer = section.footer.paragraphs[0]

    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer.add_run(
        "LegalEase - AI Generated Draft"
    )

    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(8)

    output = io.BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


def is_heading(text: str) -> bool:

    upper = text.upper().strip()

    heading_words = [
        "PARTIES",
        "EFFECTIVE DATE",
        "RECITALS",
        "BACKGROUND",
        "SIGNATURES",
        "LEGAL NOTICE",
        "TERMS AND CONDITIONS",
    ]

    if upper in heading_words:
        return True

    if re.match(r"^\d+\.\s+[A-Z]", text):
        return True

    if len(text) < 70 and text.isupper():
        return True

    return False


class LegalEasePDF(FPDF):

    def __init__(self, doc_type: str):
        super().__init__()

        self.doc_type = doc_type

        self.set_auto_page_break(
            auto=True,
            margin=18
        )

    def header(self):

        self.set_font(
            "Helvetica",
            "B",
            12
        )

        self.cell(
            0,
            8,
            "LegalEase",
            align="C"
        )

        self.ln(8)

        self.set_draw_color(
            100,
            100,
            100
        )

        self.line(
            15,
            24,
            195,
            24
        )

        self.ln(5)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            "",
            8
        )

        self.cell(
            0,
            8,
            "LegalEase - AI Generated Draft",
            align="C"
        )


def format_pdf(
    text: str,
    doc_type: str
) -> bytes:

    pdf = LegalEasePDF(doc_type)

    pdf.set_margins(
        left=18,
        top=30,
        right=18
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        "B",
        16
    )

    pdf.cell(
        0,
        10,
        doc_type.upper(),
        align="C"
    )

    pdf.ln(15)

    blocks = split_document(text)

    for block in blocks:

        if is_heading(block):

            pdf.set_font(
                "Helvetica",
                "B",
                11
            )

            pdf.multi_cell(
                0,
                7,
                block
            )

            pdf.ln(2)

        else:

            pdf.set_font(
                "Helvetica",
                "",
                10
            )

            pdf.multi_cell(
                0,
                6,
                block
            )

            pdf.ln(2)

    output = pdf.output()

    if isinstance(output, bytearray):
        output = bytes(output)

    return output


def format_txt(text: str) -> bytes:

    clean = sanitize_text(text)

    return clean.encode(
        "utf-8"
    )